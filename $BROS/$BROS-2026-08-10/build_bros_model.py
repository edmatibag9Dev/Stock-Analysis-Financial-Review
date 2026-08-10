#!/usr/bin/env python3
"""Build $BROS_Investment_Model_2026-08-10.xlsx — EBITDA-terminal DCF, unit economics,
leadership scorecard, options strategy. Single 12% WACC across scenarios."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

OUT = "$BROS_Investment_Model_2026-08-10.xlsx"

BLUE = Font(name="Arial", size=10, color="0000FF")           # hardcoded inputs
BLUE_B = Font(name="Arial", size=10, color="0000FF", bold=True)
BLACK = Font(name="Arial", size=10, color="000000")          # formulas
BLACK_B = Font(name="Arial", size=10, bold=True)
GREEN = Font(name="Arial", size=10, color="008000")          # cross-sheet links
GREEN_B = Font(name="Arial", size=10, color="008000", bold=True)
H1 = Font(name="Arial", size=14, bold=True, color="FFFFFF")
H2 = Font(name="Arial", size=11, bold=True, color="FFFFFF")
LABEL = Font(name="Arial", size=10)
LABEL_B = Font(name="Arial", size=10, bold=True)
SMALL_IT = Font(name="Arial", size=8, italic=True, color="666666")

NAVY = PatternFill("solid", fgColor="1F4E79")
GREY = PatternFill("solid", fgColor="D9D9D9")
YELLOW = PatternFill("solid", fgColor="FFFF00")
LT_RED = PatternFill("solid", fgColor="FFC7CE")
LT_YEL = PatternFill("solid", fgColor="FFEB9C")
LT_GRN = PatternFill("solid", fgColor="C6EFCE")

CUR = '$#,##0;($#,##0);"-"'
CUR2 = '$#,##0.00;($#,##0.00);"-"'
PCT = '0.0%;(0.0%);"-"'
MULT = '0.0"x"'
NUM = '#,##0;(#,##0);"-"'

thin = Side(style="thin", color="BFBFBF")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = openpyxl.Workbook()

def title(ws, cell, text, span=None):
    ws[cell] = text
    ws[cell].font = H1
    ws[cell].fill = NAVY
    if span:
        ws.merge_cells(span)

def sub(ws, cell, text, span=None):
    ws[cell] = text
    ws[cell].font = H2
    ws[cell].fill = NAVY
    if span:
        ws.merge_cells(span)

def put(ws, cell, value, font=BLACK, fmt=None, fill=None, comment=None, align=None, border=True):
    c = ws[cell]
    c.value = value
    c.font = font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    if comment: c.comment = Comment(comment, "Model")
    if align: c.alignment = Alignment(horizontal=align)
    if border: c.border = BORDER
    return c

# ============================================================ MODEL SHEET
ws = wb.active
ws.title = "BROS_Model"
ws.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGH", [34, 13, 13, 13, 13, 13, 13, 13]):
    ws.column_dimensions[col].width = w

title(ws, "A1", "DUTCH BROS ($BROS) — DCF MODEL (EBITDA-TERMINAL METHOD) — 2026-08-10", "A1:H1")
ws["A2"] = "Single 12% WACC across all scenarios. Terminal EV/EBITDA calibrated to terminal EBITDA margin. Restaurant methodology (not FCF Gordon Growth)."
ws["A2"].font = SMALL_IT

# ---- Global inputs
sub(ws, "A4", "GLOBAL INPUTS (blue = hardcoded)", "A4:B4")
gi = [
    ("FY2026 Revenue — guidance midpoint ($M)", 2115, CUR, "Source: Q2 2026 8-K (Aug 5, 2026) — FY26 guidance raised to $2.10–$2.13B; midpoint $2,115M"),
    ("FY2026 Adj. EBITDA — guidance midpoint ($M)", 387.5, CUR, "Source: Q2 2026 8-K — FY26 adj. EBITDA guidance raised to $385–$390M; midpoint $387.5M"),
    ("WACC (single rate, all scenarios)", 0.12, PCT, "Methodology rule: one WACC across bull/base/bear. 12% per restaurant/pre-FCF default"),
    ("Net cash ($M)", 74, CUR, "Q2 2026 10-Q balance sheet: $268.6M cash − $194.6M LT debt ≈ $74M net cash (6/30/26)"),
    ("Fully exchanged diluted shares (M)", 180, NUM, "Q2 2026 release: adjusted fully exchanged wtd. diluted 178.1M; 180M used for creep"),
    ("Years to terminal (FY2030 year-end)", 4.4, '0.0', "Aug 2026 → Dec 2030 ≈ 4.4 years"),
    ("Cash taxes (% of EBITDA, simplification)", 0.15, PCT, "Simplifying assumption: cash taxes ≈ 15% of EBITDA"),
    ("Capex %% of revenue — FY2026", 0.17, PCT, "FY26 capex guide $350–370M ÷ $2,115M rev ≈ 17%"),
    ("Capex %% of revenue — FY2030 (glide)", 0.10, PCT, "Assumes growth capex moderates as fleet matures"),
]
r = 5
for label, v, fmt, src in gi:
    put(ws, f"A{r}", label.replace("%%", "%"), LABEL)
    put(ws, f"B{r}", v, BLUE, fmt, comment=src)
    r += 1
REV0, EBITDA0, WACC, NETCASH, SHARES, YRS, TAXPCT, CAPX0, CAPXT = [f"$B${i}" for i in range(5, 14)]

# ---- Scenario blocks
scen_defs = [
    ("BEAR", 0.12, 0.18, 12, LT_RED,
     "SSS decelerates to 2–3%, traffic goes negative, margins compress; openings slow"),
    ("BASE", 0.185, 0.22, 18, LT_YEL,
     "18–20% revenue growth; SSS normalizes 4–5%; margins expand to 22% by FY2030"),
    ("BULL", 0.225, 0.26, 22, LT_GRN,
     "22%+ growth via 2,029-shop 2029 goal + Salad and Go conversions; 26% terminal margin"),
]
years = [2026, 2027, 2028, 2029, 2030]
row = 15
scen_price_rows = {}
for name, cagr, tmargin, tmult, fill, driver in scen_defs:
    sub(ws, f"A{row}", f"{name} CASE — {driver}", f"A{row}:H{row}")
    ws[f"A{row}"].fill = NAVY
    r0 = row + 1
    put(ws, f"A{r0}", "Revenue CAGR FY26→FY30", LABEL)
    put(ws, f"B{r0}", cagr, BLUE, PCT, fill=YELLOW, comment="Scenario lever")
    put(ws, f"A{r0+1}", "Terminal EBITDA margin (FY2030)", LABEL)
    put(ws, f"B{r0+1}", tmargin, BLUE, PCT, fill=YELLOW, comment="Scenario lever — terminal multiple calibrated to this margin")
    put(ws, f"A{r0+2}", "Terminal EV/EBITDA multiple", LABEL)
    put(ws, f"B{r0+2}", tmult, BLUE, MULT, fill=YELLOW, comment="Calibrated to terminal margin: 18%→12x, 22%→18x, 26%→22x")
    CAGRC, TMARG, TMULT = f"$B${r0}", f"$B${r0+1}", f"$B${r0+2}"
    hr = r0 + 4
    put(ws, f"A{hr}", "Projection ($M)", LABEL_B, fill=GREY)
    for i, y in enumerate(years):
        put(ws, f"{get_column_letter(3+i)}{hr}", str(y), LABEL_B, fill=GREY, align="center")
    rows = {
        "Revenue": hr+1, "EBITDA margin %": hr+2, "EBITDA": hr+3, "Capex % rev": hr+4,
        "Capex": hr+5, "Cash taxes": hr+6, "FCF": hr+7, "Discount factor": hr+8, "PV of FCF": hr+9,
    }
    for lbl, rr in rows.items():
        put(ws, f"A{rr}", lbl, LABEL)
    for i, y in enumerate(years):
        col = get_column_letter(3+i)
        if i == 0:
            put(ws, f"{col}{rows['Revenue']}", f"={REV0}", BLACK, CUR)
            put(ws, f"{col}{rows['EBITDA margin %']}", f"={EBITDA0}/{REV0}", BLACK, PCT)
        else:
            put(ws, f"{col}{rows['Revenue']}", f"={get_column_letter(2+i)}{rows['Revenue']}*(1+{CAGRC})", BLACK, CUR)
            put(ws, f"{col}{rows['EBITDA margin %']}",
                f"={EBITDA0}/{REV0}+({TMARG}-{EBITDA0}/{REV0})*{i}/4", BLACK, PCT)
        put(ws, f"{col}{rows['EBITDA']}", f"={col}{rows['Revenue']}*{col}{rows['EBITDA margin %']}", BLACK, CUR)
        put(ws, f"{col}{rows['Capex % rev']}", f"={CAPX0}+({CAPXT}-{CAPX0})*{i}/4", BLACK, PCT)
        put(ws, f"{col}{rows['Capex']}", f"={col}{rows['Revenue']}*{col}{rows['Capex % rev']}", BLACK, CUR)
        put(ws, f"{col}{rows['Cash taxes']}", f"={col}{rows['EBITDA']}*{TAXPCT}", BLACK, CUR)
        put(ws, f"{col}{rows['FCF']}", f"={col}{rows['EBITDA']}-{col}{rows['Capex']}-{col}{rows['Cash taxes']}", BLACK, CUR)
        if i == 0:
            put(ws, f"{col}{rows['Discount factor']}", 0, BLACK, '0.000', comment="FY2026 partially elapsed — excluded from PV")
        else:
            put(ws, f"{col}{rows['Discount factor']}", f"=1/(1+{WACC})^({i}+0.4)", BLACK, '0.000',
                comment="Mid-Aug 2026 valuation date: FY27 ≈ 1.4 yrs out, … FY30 ≈ 4.4 yrs")
        put(ws, f"{col}{rows['PV of FCF']}", f"={col}{rows['FCF']}*{col}{rows['Discount factor']}", BLACK, CUR)
    vr = hr + 11
    valuation = [
        ("Terminal EV (FY30 EBITDA × multiple) ($M)", f"=G{rows['EBITDA']}*{TMULT}", CUR),
        ("PV of Terminal Value ($M)", f"=G{rows['EBITDA']}*{TMULT}/(1+{WACC})^{YRS}", CUR),
        ("PV of interim FCF FY27–FY30 ($M)", f"=SUM(D{rows['PV of FCF']}:G{rows['PV of FCF']})", CUR),
        ("Enterprise Value ($M)", f"=B{vr+1}+B{vr+2}", CUR),
        ("+ Net cash ($M)", f"={NETCASH}", CUR),
        ("Equity Value ($M)", f"=B{vr+3}+B{vr+4}", CUR),
        ("÷ Fully exchanged diluted shares (M)", f"={SHARES}", NUM),
        ("INTRINSIC PRICE / SHARE", f"=B{vr+5}/B{vr+6}", CUR2),
    ]
    for j, (lbl, f, fmt) in enumerate(valuation):
        bold = (j == len(valuation) - 1)
        put(ws, f"A{vr+j}", lbl, LABEL_B if bold else LABEL, fill=fill if bold else None)
        put(ws, f"B{vr+j}", f, BLACK_B if bold else BLACK, fmt, fill=fill if bold else None)
    scen_price_rows[name] = vr + len(valuation) - 1
    row = vr + len(valuation) + 2

# ============================================================ DASHBOARD
ds = wb.create_sheet("Dashboard", 0)
ds.sheet_view.showGridLines = False
for col, w in zip("ABCD", [36, 16, 16, 16]):
    ds.column_dimensions[col].width = w
title(ds, "A1", "DUTCH BROS ($BROS) — DASHBOARD — 2026-08-10", "A1:D1")
ds["A2"] = "Quarterly update of 2026-06-05 analysis. Blue = hardcoded inputs, black = formulas, green = cross-sheet links."
ds["A2"].font = SMALL_IT

sub(ds, "A4", "MARKET DATA", "A4:B4")
mkt = [
    ("Current price (Aug 10, 2026 close)", 51.29, CUR2, "stockanalysis.com close 8/10/26; fell ~18% on 8/6 after Q2 print"),
    ("52-week high / low", "$74.65 / $44.58", None, "stockanalysis.com"),
    ("Fully exchanged diluted shares (M)", 180, NUM, "Q2'26 release adj. fully exchanged wtd. diluted 178.1M; 180M used"),
    ("Market cap ($M)", "=B5*B7", CUR, None),
    ("Net cash ($M)", 74, CUR, "Q2'26 10-Q: $268.6M cash − $194.6M LT debt"),
    ("Enterprise value ($M)", "=B8-B9", CUR, None),
    ("EV / FY2026E EBITDA ($387.5M guide mid)", "=B10/387.5", MULT, "FY26 adj. EBITDA guide $385–390M"),
]
r = 5
for lbl, v, fmt, src in mkt:
    put(ds, f"A{r}", lbl, LABEL)
    put(ds, f"B{r}", v, BLUE if src else BLACK, fmt, comment=src)
    r += 1

sub(ds, "A13", "Q2 2026 ACTUALS (8-K, Aug 5 2026)", "A13:B13")
q2 = [
    ("Revenue ($M)", 550.9, CUR, "+32.5% YoY vs $415.8M"),
    ("Revenue growth YoY", 0.325, PCT, None),
    ("Systemwide SSS", 0.058, PCT, "Co-op SSS +8.3%; system transactions +1.7%, ticket +4.1%"),
    ("Systemwide transaction growth", 0.017, PCT, "Down from +3.7% PY — the traffic deceleration signal"),
    ("Co-op shop contribution margin", 0.306, PCT, "vs 31.1% PY"),
    ("Adj. EBITDA ($M)", 113.7, CUR, "20.6% margin vs 21.4% PY"),
    ("Adj. EPS (fully exchanged)", 0.33, CUR2, "vs $0.26 PY"),
    ("Total shops (period end)", 1225, NUM, "888 co-op / 337 franchised; 48 opened in Q2, 89 in H1"),
]
r = 14
for lbl, v, fmt, src in q2:
    put(ds, f"A{r}", lbl, LABEL)
    put(ds, f"B{r}", v, BLUE, fmt, comment=src)
    r += 1

sub(ds, "A23", "FY2026 GUIDANCE (RAISED AUG 5)", "A23:B23")
gd = [
    ("Revenue", "$2.10–$2.13B (was $2.05–$2.08B)", "Excludes Salad and Go"),
    ("Systemwide SSS", "5–6%", None),
    ("Adj. EBITDA", "$385–$390M (was $370–$380M)", None),
    ("Capex", "$350–$370M", None),
    ("System shop openings", "≥185", None),
]
r = 24
for lbl, v, src in gd:
    put(ds, f"A{r}", lbl, LABEL)
    put(ds, f"B{r}", v, BLUE, comment=src)
    r += 1

sub(ds, "A30", "DCF INTRINSIC VALUE (EBITDA-TERMINAL, 12% WACC)", "A30:D30")
put(ds, "A31", "Scenario", LABEL_B, fill=GREY)
put(ds, "B31", "Bear", LABEL_B, fill=LT_RED, align="center")
put(ds, "C31", "Base", LABEL_B, fill=LT_YEL, align="center")
put(ds, "D31", "Bull", LABEL_B, fill=LT_GRN, align="center")
put(ds, "A32", "Intrinsic price / share", LABEL_B)
put(ds, "B32", f"=BROS_Model!B{scen_price_rows['BEAR']}", GREEN_B, CUR2, align="center")
put(ds, "C32", f"=BROS_Model!B{scen_price_rows['BASE']}", GREEN_B, CUR2, align="center")
put(ds, "D32", f"=BROS_Model!B{scen_price_rows['BULL']}", GREEN_B, CUR2, align="center")
put(ds, "A33", "vs current price", LABEL)
for col in "BCD":
    put(ds, f"{col}33", f"={col}32/$B$5-1", BLACK, PCT, align="center")

sub(ds, "A36", "TECHNICAL SNAPSHOT", "A36:B36")
tech = [
    ("20-day SMA", 63.28, CUR2, "Finviz, Aug 10 2026"),
    ("50-day SMA", 64.21, CUR2, "Finviz"),
    ("200-day SMA", 57.24, CUR2, "Finviz — price below all three SMAs"),
    ("RSI (14)", 20, NUM, "stockinvest.us Aug 7 — deeply oversold"),
    ("Support", "$50–51, then $44.58", None, "psych level + 52-wk low"),
    ("Resistance", "$57–58 (200d), $63–64 (20/50d)", None, "gap-fill zone $57.5–65.6"),
]
r = 37
for lbl, v, *rest in tech:
    fmt = rest[0] if rest else None
    src = rest[1] if len(rest) > 1 else None
    put(ds, f"A{r}", lbl, LABEL)
    put(ds, f"B{r}", v, BLUE, fmt, comment=src)
    r += 1

# ============================================================ UNIT ECONOMICS
ue = wb.create_sheet("Unit_Economics")
ue.sheet_view.showGridLines = False
for col, w in zip("ABCDEFG", [30, 14, 14, 14, 14, 14, 30]):
    ue.column_dimensions[col].width = w
title(ue, "A1", "UNIT ECONOMICS + PEER COMPS + LEADERSHIP SCORECARD", "A1:G1")
ue["A2"] = "Restaurant methodology — Rule of 40 replaced by unit-economics scorecard (see non-SaaS adaptations)."
ue["A2"].font = SMALL_IT

sub(ue, "A4", "UNIT ECONOMICS SCORECARD (Q2 2026)", "A4:C4")
put(ue, "A5", "Metric", LABEL_B, fill=GREY); put(ue, "B5", "Value", LABEL_B, fill=GREY); put(ue, "C5", "Watch", LABEL_B, fill=GREY)
uec = [
    ("Systemwide SSS", "+5.8%", "Q3 guide 4–5% = deceleration"),
    ("System transaction growth", "+1.7%", "ALERT if negative — was +3.7% PY"),
    ("Avg ticket growth", "+4.1%", "Pricing contribution <1% in H2'26"),
    ("Co-op contribution margin", "30.6%", "vs 31.1% PY; coffee + occupancy pressure"),
    ("Adj. EBITDA margin", "20.6%", "vs 21.4% PY"),
    ("Net new shops (Q2 / H1)", "48 / 89", "On track for ≥185 FY26"),
    ("Total shops", "1,225 of 4,000+ potential", "~31% penetration; 2,029-shop 2029 goal, ~90% pipeline secured"),
    ("Est. AUV (systemwide)", "~$2.2M (est.)", "Q1'26 was $2.16M +6.6% YoY"),
]
r = 6
for a, b, c in uec:
    put(ue, f"A{r}", a, LABEL); put(ue, f"B{r}", b, BLUE); put(ue, f"C{r}", c, LABEL)
    r += 1

sub(ue, "A16", "PEER COMPARABLES (approximate, Aug 2026)", "A16:G16")
hdrs = ["Company", "Ticker", "SSS %", "RL margin %", "EV/EBITDA (fwd)", "Rev growth (fwd)", "Note"]
for i, h in enumerate(hdrs):
    put(ue, f"{get_column_letter(1+i)}17", h, LABEL_B, fill=GREY)
peers = [
    ("Dutch Bros", "BROS", "+5.8%", "~31%", "~24x", "~25–28%", "Subject — fastest grower in set"),
    ("Starbucks", "SBUX", "~+1–3%", "n/a", "~16–18x", "~3–5%", "Scale leader, traffic-challenged"),
    ("Chipotle", "CMG", "~+2–4%", "~26–27%", "~24–27x", "~10–12%", "Margin benchmark"),
    ("CAVA", "CAVA", "~+5–8%", "~25%", "~40x+", "~20–25%", "Closest growth comp, richer multiple"),
    ("Shake Shack", "SHAK", "~+2–4%", "~21–22%", "~22–26x", "~12–15%", "Mid-growth fast casual"),
    ("Wingstop", "WING", "~flat–+3%", "n/a (franchise)", "~38–45x", "~15–20%", "Franchise model premium"),
    ("7 Brew", "Private", "n/a", "n/a", "n/a", "Rapid", "Most direct drive-thru competitor"),
]
r = 18
for p in peers:
    fill = LT_YEL if p[1] == "BROS" else None
    for i, v in enumerate(p):
        put(ue, f"{get_column_letter(1+i)}{r}", v, BLUE if i > 1 else LABEL, fill=fill)
    r += 1
put(ue, "A25", "Peer figures are approximate (Aug 2026 web sources); directional context only.", SMALL_IT, border=False)

sub(ue, "A27", "LEADERSHIP SCORECARD (Nash-style, 0/5/10 — evidence required)", "A27:G27")
put(ue, "A28", "CEO: Christine Barone (hired, since Jan 2024; ex-Starbucks SVP, ex-True Food Kitchen CEO)", LABEL, border=False)
put(ue, "A29", "Founder: Travis Boersma, Executive Chairman — 38.8% economic / 73.1% voting (Class B 10:1 super-vote)", LABEL, border=False)
put(ue, "A31", "Criterion", LABEL_B, fill=GREY); put(ue, "B31", "Score (0/5/10)", LABEL_B, fill=GREY)
ue.merge_cells("C31:G31"); put(ue, "C31", "Evidence", LABEL_B, fill=GREY)
lead = [
    ("Execution track record", 10, "Q2'26 beat; FY26 rev guide raised to $2.10–2.13B and EBITDA to $385–390M; 32.5% rev growth; 89 H1 openings on ≥185 pace"),
    ("Founder / tenure / skin in game", 5, "Founder Exec Chairman w/ 38.8% economic stake; but hired CEO tenure only 2.5 yrs; founder in programmatic sell-down"),
    ("Capital allocation", 5, "Opportunistic distressed site buys (Clutch $20M/20 sites; Salad and Go $105M/65 sites) vs heavy capex $350–370M, negative FCF, no buybacks"),
    ("Insider activity (Form 4)", 5, "No open-market buys; chair-linked entities sold 2.25M sh H1'26 @$60–64 — all 10b5-1 (routine); CEO sold 42k 10b5-1"),
    ("Governance & alignment", 5, "Dual-class: Boersma 73.1% votes vs 38.8% economics (Class B 10 votes/sh, sunset <5%); separate CEO/Chair; no restatements"),
]
r = 32
for crit, score, ev in lead:
    put(ue, f"A{r}", crit, LABEL)
    put(ue, f"B{r}", score, BLUE, NUM, fill=YELLOW, comment="Judgment input — 0/5/10")
    ue.merge_cells(f"C{r}:G{r}")
    put(ue, f"C{r}", ev, LABEL)
    r += 1
put(ue, "A37", "COMPOSITE (max 50)", LABEL_B, fill=GREY)
put(ue, "B37", "=SUM(B32:B36)", BLACK_B, NUM, fill=GREY)
ue.merge_cells("C37:G37")
put(ue, "C37", '=IF(B37>=40,"Strong — feeds Bull Case",IF(B37>=25,"Neutral — not a differentiator","Risk — feeds Bear Case"))', BLACK_B)

# ============================================================ OPTIONS
op = wb.create_sheet("Options_Strategy")
op.sheet_view.showGridLines = False
for col, w in zip("ABCDEFG", [26, 20, 12, 16, 16, 16, 40]):
    op.column_dimensions[col].width = w
title(op, "A1", "OPTIONS STRATEGY — LONG EQUITY HOLDER OVERLAY", "A1:G1")
sub(op, "A3", "MARKET DATA REFERENCE", "A3:B3")
od = [
    ("Current price", 51.29, CUR2, "Aug 10, 2026 close"),
    ("Est. IV (post-earnings)", "~45–55% (est.)", None, "Estimate — verify live quotes; daily vol ~4.2% last week"),
    ("DCF base (model)", f"=BROS_Model!B{scen_price_rows['BASE']}", CUR2, None),
    ("DCF bear (model)", f"=BROS_Model!B{scen_price_rows['BEAR']}", CUR2, None),
    ("Support / Resistance", "$50–51 / $57–58", None, None),
]
r = 4
for lbl, v, fmt, src in od:
    put(op, f"A{r}", lbl, LABEL)
    is_link = isinstance(v, str) and v.startswith("=")
    put(op, f"B{r}", v, GREEN if is_link else BLUE, fmt, comment=src)
    r += 1

sub(op, "A10", "RECOMMENDED STRATEGIES (premiums are estimates — verify live)", "A10:G10")
oh = ["Strategy", "Structure", "Expiry", "Est. premium", "Max gain", "Max loss", "Rationale"]
for i, h in enumerate(oh):
    put(op, f"{get_column_letter(1+i)}11", h, LABEL_B, fill=GREY)
strats = [
    ("Covered Call (income)", "Sell $57.5–$60 Call", "6–8 wks", "~$1.30–$1.90/sh (~2.5–3.5% per cycle)",
     "Premium + upside to strike", "Shares called away above strike (opportunity cost)",
     "Strike sits at 200-day SMA / gap-fill resistance; per-cycle yield (NOT annualized)"),
    ("Cash-Secured Put (accumulate)", "Sell $47.5 Put", "6–8 wks", "~$1.60–$2.20/sh (~3.4–4.6% per cycle)",
     "Keep premium if >$47.5 at expiry", "$4,530/contract if BROS → $0 (assigned at $47.5 − premium)",
     "Effective entry ~$45.5–46 — ~22% below rebuilt DCF base; disciplined add level"),
    ("CSP ladder (aggressive add)", "Sell $45 Put", "8–10 wks", "~$1.20–$1.70/sh (~2.7–3.8% per cycle)",
     "Keep premium if >$45", "$4,330/contract at max ($45 − premium)",
     "Assignment near 52-wk low $44.58; only if sized for bear case $26"),
]
r = 12
for s in strats:
    for i, v in enumerate(s):
        put(op, f"{get_column_letter(1+i)}{r}", v, BLUE if i in (1, 2, 3) else LABEL)
    r += 1
put(op, "A16", "Sizing rule: compute max loss in dollars before sizing. Bear-case DCF is ~$26 — size so a move there is acceptable. Yields are per-cycle, not annualized.", SMALL_IT, border=False)

wb.save(OUT)
print("saved", OUT)
